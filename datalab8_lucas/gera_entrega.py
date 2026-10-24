"""Monta a entrega do DataLab 8 no template oficial usado nos demais trabalhos."""
import json
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import sys

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "scripts"))
import entrega

entrega.ALUNO = "Lucas Fiche Ungarelli Borges"
entrega.MATRICULA = "2021068"
TEMPLATE, _preenche_tabelas = entrega.TEMPLATE, entrega._preenche_tabelas

RAIZ = Path(__file__).resolve().parent.parent
LAB = RAIZ / "datalab8_lucas"
IMAGENS = LAB / "imagens"
SAIDA = LAB / "entrega"
NOME = "IDP_INE_DataLab8-B2_Lucas_Fiche_Ungarelli_Borges"


def numero(v,casas=2):
    return f"{v:,.{casas}f}".replace(",","@").replace(".",",").replace("@",".")


def reais(v,casas=2):
    return "R$ "+numero(v,casas)


def paragrafo(doc,texto,negrito=False,tamanho=11.5):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after=Pt(8)
    p.paragraph_format.line_spacing=1.16
    p.paragraph_format.widow_control=True
    r=p.add_run(texto);r.font.name="Arial";r.font.size=Pt(tamanho);r.bold=negrito
    r.font.color.rgb=RGBColor(0,0,0)
    if negrito:p.paragraph_format.keep_with_next=True
    return p


def titulo(doc,texto):
    p=paragrafo(doc,texto,True,14)
    p.paragraph_format.space_after=Pt(14)
    return p


def imagem(doc,caminho,largura,legenda):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(caminho),width=Inches(largura))
    p=paragrafo(doc,legenda,tamanho=9)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p


def secao(doc,paisagem=False):
    s=doc.add_section(WD_SECTION_START.NEW_PAGE)
    s.orientation=WD_ORIENT.LANDSCAPE if paisagem else WD_ORIENT.PORTRAIT
    s.page_width=Inches(11.693 if paisagem else 8.268)
    s.page_height=Inches(8.268 if paisagem else 11.693)
    s.top_margin=s.bottom_margin=Inches(.65 if paisagem else 1)
    s.left_margin=s.right_margin=Inches(.65 if paisagem else 1)
    s.header_distance=Inches(.25 if paisagem else .5)
    s.footer_distance=Inches(.25 if paisagem else .5)
    s.header.is_linked_to_previous=False
    for p in s.header.paragraphs:p.clear()
    s.footer.is_linked_to_previous=False
    for p in s.footer.paragraphs:p.clear()
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    from docx.oxml import OxmlElement
    fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");p._p.append(fld)
    return s


def main():
    s=json.loads((LAB/"dados/indicadores.json").read_text())
    doc=Document(TEMPLATE)
    _preenche_tabelas(doc,"DataLab 8-B2 - Mercado de trabalho em TI no Brasil")
    inicio=next(p._p for p in doc.paragraphs if p.text.startswith("Solução do DevLab"))
    body=doc._element.body
    apagando=False
    for elemento in list(body):
        if elemento is inicio:apagando=True
        if apagando and elemento.tag!=qn("w:sectPr"):body.remove(elemento)
    doc.core_properties.title="DataLab 8 B2 Mercado de trabalho em TI no Brasil"
    doc.core_properties.author="Lucas Fiche Ungarelli Borges"
    doc.core_properties.last_modified_by="Lucas Fiche Ungarelli Borges"
    doc.core_properties.subject="Inteligência de Negócio - APINFO 2025 - Brasil e Região Sul"

    titulo(doc,"1. Solução do DataLab")
    paragrafo(doc,"Esta análise usa a pesquisa salarial APINFO 2025 para estudar a remuneração em TI no Brasil e nos três estados da Região Sul. O dataset reúne as tabelas de Brasil, Paraná, Rio Grande do Sul e Santa Catarina. A comparação regional se concentra em arquitetura de sistemas e desenvolvimento de software, mantendo cargo e vínculo iguais em cada comparação.")
    imagem(doc,IMAGENS/"00_tratamento_dataset.png",6.2,
           "Figura 1 - Valores da tabela Paraná antes e depois do tratamento do dataset.")
    paragrafo(doc,"A base tem 664 registros: 83 perfis profissionais, quatro localidades e dois vínculos. Há 656 valores publicados e oito ausentes. Os campos vazios foram preservados como nulos e o estágio foi identificado como bolsa, embora a fonte o apresente na coluna CLT.")
    paragrafo(doc,"As cinco páginas seguintes respondem ao enunciado. O dashboard foi construído em Python, com exportação para PDF e HTML navegável. Os dados brutos, o CSV tratado, os cálculos e os scripts acompanham a entrega e permitem refazer os resultados.")

    secao(doc,True)
    nomes=["01_media_brasil","02_cargo_destaque","03_clt_pj","04_extremos","05_comparativo_estados"]
    legendas=["A - Média das referências nacionais e comparação das médias por localidade.","B - Cargo com maior remuneração na tabela nacional.","C - Ganho percentual PJ/CLT e diferenças nos cargos de software selecionados.","D - Maiores e menores referências nacionais e estaduais, com identificação das bolsas.","E - Arquitetura e desenvolvimento de software em PR, RS e SC comparados ao Brasil."]
    for i,(nome,legenda) in enumerate(zip(nomes,legendas)):
        if i:doc.add_page_break()
        imagem(doc,IMAGENS/(nome+".png"),10.35,f"Figura {i+2} - Página {i+1} do dashboard. {legenda}")

    secao(doc)
    titulo(doc,"2. Descrição técnica do laboratório")
    paragrafo(doc,"A pesquisa APINFO foi realizada pela internet entre agosto e setembro de 2025, com 26.174 participantes no levantamento completo e maior concentração no Sudeste e no Sul. As tabelas foram consultadas em 24/10/2026. Todos os rendimentos analisados pertencem à mesma edição e estão em reais nominais por mês, sem correção pela inflação.")
    paragrafo(doc,"A. Qual foi o objetivo do DataLab?",True)
    paragrafo(doc,"Organizar os dados públicos da APINFO e responder às questões sobre remuneração média, cargos, contratação e diferenças regionais. A escolha de Paraná, Rio Grande do Sul e Santa Catarina permite observar os três estados de uma mesma região. Para a comparação de cargos, foram selecionados Arquiteto de Sistemas e Desenvolvedores Full-Stack, Front-End e Back-End, presentes nas quatro tabelas.")
    paragrafo(doc,"B. Quais foram as técnicas utilizadas?",True)
    paragrafo(doc,"A extração identifica a tabela com as colunas Cargo, Especialidade, CLT e PJ. Os arquivos HTML são preservados, juntamente com endereço de origem, data de coleta e hash SHA-256. O tratamento transforma CLT e PJ em linhas, com uma chave formada por ano, UF, perfil e vínculo. A tabela Brasil permanece separada dos estados para evitar duplicar referências no cálculo nacional.")
    paragrafo(doc,"Foram removidos espaços excedentes e pontos de milhar para converter os valores em números. As grafias Phyton e Java Script foram normalizadas para Python e JavaScript, mantendo o conteúdo original em outras colunas. Campos sem valor não recebem zero. Estagiário é classificado como bolsa_estagio, permitindo mostrar os extremos e as médias com essa distinção.")
    paragrafo(doc,"A média nacional usa os 164 valores válidos da tabela Brasil, com igual peso por referência. O ranking considera a média de CLT e PJ dos 81 perfis completos. O ganho médio é calculado em duas etapas: primeiro (PJ / CLT - 1) × 100 para cada perfil; depois, a média dos 81 percentuais. Dividir a média PJ pela média CLT produz uma medida diferente, apresentada apenas como complemento.")
    paragrafo(doc,"A análise regional reúne 24 comparações: quatro cargos, dois vínculos e três estados. Cada valor estadual é confrontado com o valor Brasil do mesmo cargo e vínculo por meio de (estado / Brasil - 1) × 100. As tabelas não fornecem número de participantes por célula. Assim, os cálculos descrevem referências agregadas, sem estimar a média ponderada dos trabalhadores ou a renda líquida de cada regime.")
    paragrafo(doc,"A validação verifica campos ausentes, valores positivos, chaves únicas, ano das fontes e correspondência entre perfis. Também confere se mudanças nos valores estaduais deixam os indicadores nacionais inalterados. Os gráficos usam o CSV tratado e os indicadores calculados pelo script, mantendo o vínculo entre dados e apresentação.")

    doc.add_page_break()
    titulo(doc,"Resultados e aprendizado")
    paragrafo(doc,"C. O que você aprendeu?",True)
    paragrafo(doc,f"A média das referências nacionais é {reais(s['media_referencias_brasil'])}; sem a bolsa de estágio, passa a {reais(s['media_sem_bolsa_estagio'])}. As médias simples das tabelas estaduais são R$ 7.554,92 no Paraná, R$ 7.709,10 no Rio Grande do Sul e R$ 8.064,95 em Santa Catarina. Como todas usam 164 valores publicados, essa comparação mantém o mesmo número de referências, mas não representa o salário médio dos trabalhadores de cada estado.")
    paragrafo(doc,f"Gerente de TI lidera a remuneração nacional: {reais(s['lider']['CLT'],0)} na coluna CLT e {reais(s['lider']['PJ'],0)} na PJ. A média entre os vínculos é {reais(s['lider']['media_perfil_brl'],0)}. O ganho percentual médio PJ/CLT de todos os perfis completos é {numero(s['ganho_pj_medio_pct'])}%. Nos cargos escolhidos para a comparação regional, Desenvolvedor Full-Stack apresenta a maior diferença nacional, de 42,23%.")
    paragrafo(doc,f"Na tabela Brasil, o maior valor publicado é {reais(s['maior_referencia']['salario_mensal_brl'],0)}, de Gerente de TI PJ, e o menor é {reais(s['menor_referencia']['salario_mensal_brl'],0)}, da bolsa de Estagiário. Ao excluir estágio, o menor salário é {reais(s['menor_sem_bolsa']['salario_mensal_brl'],0)}, de Assistente de Informática CLT. Na base completa, incluindo os estados, o máximo continua sendo o da tabela Brasil, enquanto o mínimo cai para {reais(s['menor_base_coletada']['salario_mensal_brl'],0)}, bolsa de Estagiário no Paraná.")
    paragrafo(doc,"As 24 referências estaduais dos quatro cargos ficam abaixo do valor Brasil correspondente. Santa Catarina tem o maior valor regional em seis dos oito pares de cargo e vínculo. O Rio Grande do Sul lidera em Arquiteto de Sistemas PJ e Desenvolvedor Full-Stack CLT. Isso mostra que a posição dos estados muda conforme o perfil e o regime de contratação.")
    paragrafo(doc,"A maior distância da referência nacional aparece em Desenvolvedor Full-Stack PJ no Paraná: R$ 11.190 contra R$ 13.701 no Brasil, diferença de -18,33%. O menor afastamento é de Desenvolvedor Back-End CLT em Santa Catarina: R$ 6.950 contra R$ 7.201, ou -3,49%. Esses resultados valem para os cargos selecionados e não permitem generalizar a diferença para todo o mercado de trabalho estadual.")
    paragrafo(doc,"O exercício evidencia a importância de comparar unidades equivalentes. Trocar cargo, vínculo ou período altera o significado do resultado. Também é necessário separar bolsa de salário e lembrar que maior valor nominal de PJ não significa maior renda líquida, pois benefícios, impostos, férias e intervalos sem faturamento não foram equalizados.")
    paragrafo(doc,"Fontes dos dados",True)
    for nome,sufixo in [("Pesquisa e metodologia","p25sumario"),("Brasil","p25sal-br"),("Paraná","p25sal-pr"),("Rio Grande do Sul","p25sal-rs"),("Santa Catarina","p25sal-sc")]:
        paragrafo(doc,f"APINFO 2025 - {nome}: https://www.apinfo2.com/apinfo/informacao/{sufixo}.cfm",tamanho=9)
    SAIDA.mkdir(parents=True,exist_ok=True)
    destino=SAIDA/(NOME+".docx")
    doc.save(destino)
    print(destino.relative_to(RAIZ))
    with tempfile.TemporaryDirectory(prefix="datalab8-lucas-lo-") as pasta:
        perfil=(Path(pasta)/"perfil").as_uri()
        subprocess.run(["soffice",f"-env:UserInstallation={perfil}","--headless","--convert-to","pdf","--outdir",str(SAIDA),str(destino)],check=True,capture_output=True)
    if not destino.with_suffix(".pdf").exists():
        raise RuntimeError("O LibreOffice não gerou o PDF da entrega")
    print(destino.with_suffix(".pdf").relative_to(RAIZ))


if __name__=="__main__":
    main()
