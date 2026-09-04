"""
Montagem das entregas das Atividades Avaliativas de Inteligência de Negócio (IDP).

Preenche o "[Template aluno]" oficial da disciplina (docs/) com a identificação do
aluno, os prints das análises e as respostas da descrição técnica, e converte o
resultado para PDF com o LibreOffice.
"""
import copy
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
TEMPLATE = RAIZ / "docs" / (
    "[Template aluno] IDP.Inteligência de Negócio - "
    "Atividade Avaliativa Nr.XX (Prof. Bruno Miranda, 2026.2S).docx"
)

ALUNO = "Claudio da Aparecida Meireles Filho"
MATRICULA = "2321070"
LARGURA_IMAGEM = Inches(6.2)  # A4 com margens de 1 polegada

MARCADOR_PRINT = "[Cole aqui o print que evidencie o seu trabalho.]"
MARCADOR_DESCRICAO = "[Descreva tecnicamente"


PRETO = RGBColor(0x00, 0x00, 0x00)


def _texto(paragrafo, texto):
    """Troca o conteúdo do parágrafo mantendo a formatação do primeiro run."""
    if not paragrafo.runs:
        paragrafo.add_run("")
    paragrafo.runs[0].text = texto
    for run in paragrafo.runs[1:]:
        run._element.getparent().remove(run._element)
    return paragrafo


def _normaliza(paragrafo):
    """Tira o visual de placeholder do template (azul e itálico) do texto preenchido."""
    for run in paragrafo.runs:
        run.italic = False
        run.font.color.rgb = PRETO
    return paragrafo


def _sem_lista(paragrafo):
    """Remove a numeração herdada quando o parágrafo clonado vem de um item de lista."""
    formato = paragrafo._element.find(qn("w:pPr"))
    if formato is not None:
        numeracao = formato.find(qn("w:numPr"))
        if numeracao is not None:
            formato.remove(numeracao)
    return paragrafo


def _clona_depois(paragrafo, texto=""):
    """Duplica um parágrafo logo abaixo dele, herdando estilo e espaçamento."""
    novo = copy.deepcopy(paragrafo._element)
    paragrafo._element.addnext(novo)
    from docx.text.paragraph import Paragraph
    return _texto(Paragraph(novo, paragrafo._parent), texto)


def _preenche_tabelas(documento, tipo_atividade):
    rotulos = {
        "Disciplina:": "Inteligência de Negócio (INE)",
        "Professor:": "Bruno Miranda",
        "Tipo de atividade:": tipo_atividade,
        "Semestre": "2° semestre de 2026",
        "Departamento / Curso:": "Ciência da Computação / Engenharia de Software",
        "Nome do aluno:": ALUNO,
        "Matrícula:": MATRICULA,
    }
    for tabela in documento.tables:
        for linha in tabela.rows:
            rotulo = linha.cells[0].text.strip()
            if rotulo in rotulos:
                _normaliza(_texto(linha.cells[1].paragraphs[0], rotulos[rotulo]))


def _acha(documento, inicio):
    for paragrafo in documento.paragraphs:
        if paragrafo.text.strip().startswith(inicio):
            return paragrafo
    raise SystemExit(f"marcador não encontrado no template: {inicio!r}")


def _insere_prints(documento, prints):
    """Substitui o marcador de print pelas imagens, cada uma com sua legenda."""
    ancora = _acha(documento, MARCADOR_PRINT)
    _texto(ancora, "")

    atual = ancora
    for legenda, imagem in prints:
        atual = _clona_depois(atual, "")
        atual.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in atual.runs:
            run._element.getparent().remove(run._element)
        atual.add_run().add_picture(str(imagem), width=LARGURA_IMAGEM)

        atual = _clona_depois(atual, legenda)
        _sem_lista(atual)
        atual.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in atual.runs:
            run.italic = True
            run.bold = False
            run.font.size = Pt(9)
            run.font.color.rgb = PRETO


def _responde(documento, pergunta, paragrafos):
    """Escreve a resposta nas linhas em branco que seguem cada pergunta do template."""
    ancora = _acha(documento, pergunta)
    atual = ancora
    for texto in paragrafos:
        atual = _normaliza(_sem_lista(_clona_depois(atual, texto)))
        atual.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # o item de lista original tinha recuo deslocado; o texto corrido não precisa dele
        atual.paragraph_format.first_line_indent = Pt(0)
        for run in atual.runs:
            run.bold = False


def _para_pdf(caminho_docx, caminho_pdf):
    with tempfile.TemporaryDirectory() as temporario:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf",
             "--outdir", temporario, str(caminho_docx)],
            check=True, capture_output=True,
        )
        gerado = Path(temporario) / (caminho_docx.stem + ".pdf")
        if not gerado.exists():
            raise SystemExit(f"LibreOffice não gerou {gerado.name}")
        shutil.move(str(gerado), caminho_pdf)


def gerar(destino_docx, destino_pdf, tipo_atividade, introducao, prints, objetivo, tecnicas, aprendizado):
    """Gera o .docx e o .pdf da entrega a partir do template oficial."""
    documento = Document(TEMPLATE)
    _preenche_tabelas(documento, tipo_atividade)
    _insere_prints(documento, prints)
    _normaliza(_texto(_acha(documento, MARCADOR_DESCRICAO), introducao))
    _responde(documento, "Qual foi o objetivo do DevLab?", objetivo)
    _responde(documento, "Quais foram as técnicas utilizadas?", tecnicas)
    _responde(documento, "O que você aprendeu?", aprendizado)

    destino_docx.parent.mkdir(parents=True, exist_ok=True)
    documento.save(destino_docx)
    _para_pdf(destino_docx, destino_pdf)
    print(f"  {destino_docx.relative_to(RAIZ)}")
    print(f"  {destino_pdf.relative_to(RAIZ)}")
