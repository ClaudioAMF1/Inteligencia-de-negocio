"""Monta a entrega do DataLab 8 no template oficial usado nos demais trabalhos."""
import json
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from entrega import TEMPLATE, _preenche_tabelas

RAIZ = Path(__file__).resolve().parent.parent
LAB = RAIZ / "datalab8"
IMAGENS = LAB / "imagens"
SAIDA = LAB / "entrega"
NOME = "IDP_INE_DataLab8-B2_Claudio_Meireles"


def numero(v,casas=2):
    return f"{v:,.{casas}f}".replace(",","@").replace(".",",").replace("@",".")


def reais(v,casas=2):
    return "R$ "+numero(v,casas)


def paragrafo(doc,texto,negrito=False,tamanho=11.5):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after=Pt(8)
    p.paragraph_format.line_spacing=1.16
    p.paragraph_format.widow_control=True
    r=p.add_run(texto);r.font.name="Arial";r.font.size=Pt(tamanho);r.bold=negrito
    r.font.color.rgb=RGBColor(0,0,0)
    if negrito:p.paragraph_format.keep_with_next=True
    return p


def titulo(doc,texto):
    p=paragrafo(doc,texto,True,14)
    p.paragraph_format.space_after=Pt(14)
    return p


def imagem(doc,caminho,largura,legenda):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(caminho),width=Inches(largura))
    p=paragrafo(doc,legenda,tamanho=9)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p


def secao(doc,paisagem=False):
    s=doc.add_section(WD_SECTION_START.NEW_PAGE)
    s.orientation=WD_ORIENT.LANDSCAPE if paisagem else WD_ORIENT.PORTRAIT
    s.page_width=Inches(11.693 if paisagem else 8.268)
    s.page_height=Inches(8.268 if paisagem else 11.693)
    s.top_margin=s.bottom_margin=Inches(.65 if paisagem else 1)
    s.left_margin=s.right_margin=Inches(.65 if paisagem else 1)
    s.header_distance=Inches(.25 if paisagem else .5)
    s.footer_distance=Inches(.25 if paisagem else .5)
    s.header.is_linked_to_previous=False
    for p in s.header.paragraphs:p.clear()
    s.footer.is_linked_to_previous=False
    for p in s.footer.paragraphs:p.clear()
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    from docx.oxml import OxmlElement
    fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");p._p.append(fld)
    return s


def main():
    s=json.loads((LAB/"dados/indicadores.json").read_text())
    doc=Document(TEMPLATE)
    _preenche_tabelas(doc,"DataLab 8-B2 - Mercado de trabalho em TI no Brasil")
    inicio=next(p._p for p in doc.paragraphs if p.text.startswith("Solução do DevLab"))
    body=doc._element.body
    apagando=False
    for elemento in list(body):
        if elemento is inicio:apagando=True
        if apagando and elemento.tag!=qn("w:sectPr"):body.remove(elemento)
    doc.core_properties.title="DataLab 8 B2 Mercado de trabalho em TI no Brasil"
    doc.core_properties.author="Claudio da Aparecida Meireles Filho"
    doc.core_properties.subject="Inteligência de Negócio - APINFO 2025"

    titulo(doc,"1. Solução do DataLab")
    paragrafo(doc,"O trabalho analisa a pesquisa salarial APINFO 2025. O dataset foi construído com as tabelas do Brasil, de São Paulo, do Rio de Janeiro e de Minas Gerais. As cinco páginas seguintes respondem às perguntas do enunciado, com valores mensais nominais, nome e matrícula em todas elas.")
    imagem(doc,IMAGENS/"00_tratamento_dataset.png",6.2,
           "Figura 1 - Evidência do tratamento: recortes reais do HTML de origem e do CSV gerado.")
    paragrafo(doc,"A base final tem 664 linhas: 83 perfis profissionais, quatro localidades e dois vínculos. São 656 valores publicados e oito ausentes. Os campos vazios foram mantidos como nulos. Nenhum salário foi estimado ou preenchido com zero.")
    paragrafo(doc,"Os HTMLs originais, o CSV tratado, os indicadores e os scripts acompanham a entrega. O dashboard foi gerado em Python e exportado em PDF, com uma versão HTML para navegar pelas páginas. As imagens são saídas do código executado.")

    secao(doc,True)
    nomes=["01_media_brasil","02_cargo_destaque","03_clt_pj","04_extremos","05_comparativo_estados"]
    legendas=["A - Média simples das referências nacionais publicadas.","B - Remuneração por cargo e vínculo na tabela Brasil.","C - Ganho percentual PJ/CLT com comparação entre perfis equivalentes.","D - Extremos publicados na tabela Brasil, com a bolsa de estágio identificada.","E - Salários nominais de quatro cargos em SP, RJ e MG comparados ao Brasil."]
    for i,(nome,legenda) in enumerate(zip(nomes,legendas)):
        if i:doc.add_page_break()
        imagem(doc,IMAGENS/(nome+".png"),10.35,f"Figura {i+2} - Página {i+1} do dashboard. {legenda}")

    secao(doc)
    titulo(doc,"2. Descrição técnica do laboratório")
    paragrafo(doc,"Usei a edição de 2025 da pesquisa APINFO, realizada pela internet entre agosto e setembro daquele ano e com 26.174 participantes. A própria fonte informa maior concentração no Sudeste e no Sul. A consulta às tabelas foi feita em 24/10/2026. O ano da consulta não altera o período dos salários analisados.")
    paragrafo(doc,"A. Qual foi o objetivo do DataLab?",True)
    paragrafo(doc,"Construir uma base de dados a partir das tabelas da APINFO e transformá-la em um dashboard que permita comparar remuneração nacional, cargos, vínculos e estados. Escolhi São Paulo, Rio de Janeiro e Minas Gerais porque as três tabelas publicam os mesmos perfis da comparação, permitindo manter cargo e vínculo constantes.")
    paragrafo(doc,"B. Quais foram as técnicas utilizadas?",True)
    paragrafo(doc,"Extração e organização: o script localiza a tabela com as colunas Cargo, Especialidade, CLT e PJ, preserva os arquivos HTML e reorganiza os dados no formato longo. Cada linha representa um perfil, uma localidade e um vínculo em 2025. A chave considera ano, UF, perfil e vínculo. Brasil é uma referência nacional e nunca é somado aos estados.")
    paragrafo(doc,"Tratamento: removi espaços excedentes, converti o ponto de milhar para número e mantive campos ausentes como nulos. Corrigi somente as grafias Phyton para Python e Java Script para JavaScript, preservando os textos originais em outras colunas. A linha Estagiário recebeu a natureza bolsa_estagio, mesmo quando publicada na coluna CLT da fonte.")
    paragrafo(doc,"Agregações: a média nacional é a soma dos 164 valores válidos da tabela Brasil dividida por 164. Ela inclui a bolsa publicada, com uma segunda média que a exclui. O ranking de cargos usa a média simples entre CLT e PJ para cada perfil com ambos os valores. Para a diferença de contratação, calculei (PJ / CLT - 1) × 100 por perfil e depois a média dos 81 percentuais.")
    paragrafo(doc,"Comparação regional: selecionei Gerente de TI, Cientista de dados, Analista Cibersegurança e Engenheiro de dados. Para cada cargo e vínculo, calculei (valor do estado / valor do Brasil - 1) × 100. O resultado reúne 24 comparações, sem misturar CLT com PJ ou cargos diferentes.")
    paragrafo(doc,"Validação: conferi a quantidade de linhas, a unicidade das chaves, os campos ausentes e a positividade dos valores. Os resultados do dashboard são recalculados a partir do CSV. A tabela pública não informa quantos respondentes compõem cada perfil e vínculo; por isso, não apliquei pesos nem estimei média salarial dos trabalhadores.")

    doc.add_page_break()
    titulo(doc,"Resultados e aprendizado")
    paragrafo(doc,"C. O que você aprendeu?",True)
    paragrafo(doc,f"A média simples da base nacional ficou em {reais(s['media_referencias_brasil'])}. Excluindo a bolsa de estágio, ela passa para {reais(s['media_sem_bolsa_estagio'])}. A diferença mostra por que a definição da base precisa acompanhar o indicador: o valor representa referências publicadas, e não uma amostra de salários individuais disponível para recalcular a média dos participantes.")
    paragrafo(doc,f"Gerente de TI é o cargo de maior remuneração nos dois vínculos: {reais(s['lider']['CLT'],0)} na coluna CLT e {reais(s['lider']['PJ'],0)} na PJ. A média simples entre os dois valores é {reais(s['lider']['media_perfil_brl'],0)}. Cientista de dados vem em seguida no ranking pareado, com {reais(s['top10'][1]['media_perfil_brl'])}.")
    paragrafo(doc,f"O ganho percentual médio do PJ sobre o CLT foi de {numero(s['ganho_pj_medio_pct'])}%, calculado sobre 81 pares. A variação percentual entre a média PJ e a média CLT dos mesmos perfis é {numero(s['razao_das_medias_pares_pct'])}%, que é outra medida. A comparação descreve os valores mensais publicados e não permite concluir qual contratação deixa mais renda líquida, pois não equaliza benefícios, impostos ou períodos sem faturamento.")
    paragrafo(doc,f"Na tabela Brasil, o maior valor é {reais(s['maior_referencia']['salario_mensal_brl'],0)}, de Gerente de TI PJ. O menor é {reais(s['menor_referencia']['salario_mensal_brl'],0)}, da bolsa de Estagiário. Sem estágio, o menor é {reais(s['menor_sem_bolsa']['salario_mensal_brl'],0)}, de Assistente de Informática CLT. Considerando também os três estados coletados, os extremos são {reais(s['maior_base_coletada']['salario_mensal_brl'],0)} em SP, para Gerente de TI PJ, e {reais(s['menor_base_coletada']['salario_mensal_brl'],0)} em MG, para a bolsa de Estagiário. Não são máximos e mínimos de salários individuais.")
    paragrafo(doc,"Nos quatro cargos e dois vínculos escolhidos, São Paulo e Rio de Janeiro aparecem acima da referência nacional, enquanto Minas Gerais fica abaixo. Por exemplo, Gerente de TI PJ tem R$ 22.185 em SP, R$ 21.391 no RJ e R$ 19.223 em MG, diante de R$ 21.007 no Brasil. Essa leitura vale para os perfis selecionados e não deve ser generalizada para todos os trabalhadores de cada estado.")
    paragrafo(doc,"O principal aprendizado foi controlar a unidade de análise antes de calcular os indicadores. Perfil profissional não é quantidade de pessoas, ausência de valor não é salário zero e média de percentuais não é a mesma coisa que percentual entre médias. Essas diferenças mudam a interpretação do dashboard.")
    paragrafo(doc,"Fontes dos dados",True)
    for nome,sufixo in [("Pesquisa e metodologia","p25sumario"),("Brasil","p25sal-br"),("São Paulo","p25sal-sp"),("Rio de Janeiro","p25sal-rj"),("Minas Gerais","p25sal-mg")]:
        paragrafo(doc,f"APINFO 2025 - {nome}: https://www.apinfo2.com/apinfo/informacao/{sufixo}.cfm",tamanho=9)
    SAIDA.mkdir(parents=True,exist_ok=True)
    destino=SAIDA/(NOME+".docx")
    doc.save(destino)
    print(destino.relative_to(RAIZ))
    with tempfile.TemporaryDirectory(prefix="datalab8-lo-") as pasta:
        perfil=(Path(pasta)/"perfil").as_uri()
        subprocess.run(["soffice",f"-env:UserInstallation={perfil}","--headless","--convert-to","pdf", "--outdir",str(SAIDA),str(destino)],check=True,capture_output=True)
    if not destino.with_suffix(".pdf").exists():
        raise RuntimeError("O LibreOffice não gerou o PDF da entrega")
    print(destino.with_suffix(".pdf").relative_to(RAIZ))


if __name__=="__main__":
    main()
