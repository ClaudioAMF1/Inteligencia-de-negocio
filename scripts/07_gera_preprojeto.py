"""
Avaliação Continuada (Pré-projeto) - gera a entrega em .docx e .pdf.

Preenche o "Template IDP.INE - Modelo de Projeto de Inteligência de Negócio"
(docs/) com os itens que o enunciado pede: a dupla, a empresa abordada, a área de
negócio e o tema, as perguntas negociais que o estudo se propõe a resolver e as
bases de dados com o link público de cada fonte.
"""
import copy
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from entrega import PRETO, RAIZ, _clona_depois, _para_pdf, _texto

TEMPLATE = RAIZ / "docs" / (
    "Template IDP.INE - Modelo de Projeto de Inteligência de Negócio "
    "(Prof. Bruno Miranda, 2026.2S).docx"
)
SAIDA = RAIZ / "preprojeto" / "entrega"
ARQUIVO = "IDP.INE - Pré-projeto (Avaliação Continuada) - Claudio Meireles e Felipe Dutra"

TITULO = ("Precificação de combustíveis no Brasil: evolução, diferenças regionais "
          "e dispersão de preço na revenda")
DISCIPLINA = "Inteligência de Negócio (INE) - Prof. Bruno Miranda - 2º semestre de 2026"
AUTORES = ("Claudio da Aparecida Meireles Filho (2321070) e "
           "Felipe Pereira Dutra (2321017)")

EMPRESA = [
    "O estudo é feito sobre o segmento de revenda de combustíveis no Brasil, sob a ótica "
    "das distribuidoras que disputam a bomba. As empresas observadas são as que aparecem "
    "nomeadas no campo Bandeira do levantamento de preços da ANP: as maiores distribuidoras "
    "do país - Vibra Energia (bandeira BR), Ipiranga e Raízen (bandeira Shell) - e o conjunto "
    "dos postos de bandeira branca, que não têm contrato de exclusividade com nenhuma delas.",

    "A escolha se justifica porque a ANP publica o preço coletado posto a posto com a bandeira "
    "de cada revendedor. Isso permite analisar a estratégia de preço de empresas reais, com "
    "nome e marca, sem depender de dado proprietário: a base pública já traz a informação que "
    "uma área de pricing usaria para monitorar a concorrência.",

    "O recorte é nacional, com abertura para o Distrito Federal - mercado onde a dupla mora e "
    "cujo comportamento de preço pretende comparar com o das demais unidades da federação.",
]

AREA_E_TEMA = [
    ("Área de negócio: ",
     "Precificação e Inteligência de Mercado (pricing). É a área que define o preço de bomba, "
     "acompanha o preço praticado pela concorrência em cada praça e monitora a posição "
     "competitiva da rede."),

    ("Tema: ",
     "Comportamento do preço de revenda de combustíveis no Brasil - gasolina comum, etanol "
     "hidratado e diesel S-10. O trabalho vai medir como o preço evoluiu ao longo do tempo, "
     "o quanto ele varia entre regiões e unidades da federação, qual é a dispersão de preço "
     "dentro de uma mesma cidade, em que condições o etanol se torna vantajoso frente à "
     "gasolina e se a bandeira do posto se traduz em diferença consistente de preço."),

    ("Por que esse tema: ",
     "combustível é um preço que todo consumidor acompanha e que toda rede de postos precisa "
     "decidir toda semana. A coleta da ANP é semanal, posto a posto, e está publicada desde "
     "2004, o que dá ao projeto uma base com as três dimensões que o Power BI explora bem: "
     "geografia (região, UF e município), tempo (data da coleta) e produto/bandeira."),
]

PERGUNTAS = [
    "Como evoluiu o preço médio de revenda da gasolina comum, do etanol hidratado e do "
    "diesel S-10 no Brasil na última década, e quanto dessa alta se sustenta depois de "
    "descontada a inflação medida pelo IPCA?",

    "Quanto o preço do mesmo produto varia entre as regiões e as unidades da federação? "
    "Qual é a distância entre a UF mais cara e a mais barata em cada ano, e em que posição "
    "desse ranking está o Distrito Federal?",

    "Qual é a dispersão de preço dentro de um mesmo município na mesma semana de coleta? "
    "Em outras palavras: quanto o consumidor economiza, em reais por litro, saindo do posto "
    "mais caro para o mais barato da sua cidade?",

    "Em quais unidades da federação e em quais períodos o etanol hidratado foi vantajoso "
    "frente à gasolina comum, usando a regra de paridade de 70% entre os dois preços?",

    "Existe diferença consistente de preço entre as bandeiras (BR, Ipiranga, Shell) e os "
    "postos de bandeira branca? A diferença é a mesma em todas as regiões do país?",
]

INTRO_BASES = (
    "As três bases são públicas, de download livre e sem cadastro. A base principal é o "
    "Levantamento de Preços de Combustíveis da ANP, com a coleta semanal do preço praticado "
    "em postos revendedores por região, UF, município, produto, data da coleta, valor de "
    "venda e bandeira, publicada desde 2004. As outras duas entram como apoio: as vendas de "
    "derivados dão o volume mensal por UF e produto, usado para ponderar as médias de preço "
    "pelo tamanho de cada mercado; o IPCA entra como deflator, para separar alta real de "
    "combustível de inflação geral. Os arquivos da ANP também estão espelhados no Portal "
    "Brasileiro de Dados Abertos (dados.gov.br)."
)

BASES = [
    ("ANP - Série histórica de preços de combustíveis (Levantamento de Preços). "
     "Base principal.",
     "https://www.gov.br/anp/pt-br/centrais-de-conteudo/dados-abertos/serie-historica-de-precos-de-combustiveis"),

    ("ANP - Vendas de derivados de petróleo e etanol. Base de apoio (volume por UF e produto).",
     "https://www.gov.br/anp/pt-br/centrais-de-conteudo/dados-abertos/vendas-de-derivados-de-petroleo-e-etanol"),

    ("IBGE / SIDRA - IPCA, tabela 1737. Base de apoio (deflator dos preços).",
     "https://sidra.ibge.gov.br/tabela/1737"),
]


def _troca(paragrafo, marcador, texto):
    """Escreve no run que carrega o placeholder, preservando o tamanho e o peso da fonte.
    A cor azul dos placeholders do template não é mantida: o texto preenchido vai em preto."""
    runs = paragrafo.runs
    posicao = next((i for i, run in enumerate(runs) if marcador in run.text), None)
    if posicao is None:
        raise SystemExit(f"placeholder {marcador!r} não encontrado no template")
    for i, run in enumerate(runs):
        run.text = texto if i == posicao else ""
        run.font.color.rgb = PRETO
    return paragrafo


def _acha_indice(documento, inicio):
    for indice, paragrafo in enumerate(documento.paragraphs):
        if paragrafo.text.strip().startswith(inicio):
            return indice
    raise SystemExit(f"seção não encontrada no template: {inicio!r}")


def _preenche_secao(documento, cabecalho, paragrafos):
    """Troca o texto de exemplo que vem logo abaixo do cabeçalho da seção."""
    ancora = documento.paragraphs[_acha_indice(documento, cabecalho) + 1]
    atual = _texto(ancora, "")
    for item in paragrafos:
        rotulo, conteudo = item if isinstance(item, tuple) else ("", item)
        atual = _clona_depois(atual, rotulo + conteudo)
        atual.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in atual.runs:
            run.bold = False
        if rotulo:
            # separa o rótulo em um run próprio para destacá-lo em negrito
            corpo = atual.runs[0]
            corpo.text = conteudo
            destaque = copy.deepcopy(corpo._element)
            corpo._element.addprevious(destaque)
            atual.runs[0].text = rotulo
            atual.runs[0].bold = True
    ancora._element.getparent().remove(ancora._element)


def _formato(run, sem_negrito=False):
    """Cópia da formatação de fonte de um run, opcionalmente sem o negrito."""
    formato = run._element.find(qn("w:rPr"))
    if formato is None:
        return None
    copia = copy.deepcopy(formato)
    if sem_negrito:
        for negrito in copia.findall(qn("w:b")):
            copia.remove(negrito)
    return copia


def _formato_do_cabecalho(tabela):
    """Os parágrafos vazios do template não têm fonte definida e caem no padrão serifado;
    herdo a do cabeçalho da própria tabela, sem o negrito."""
    return _formato(tabela.rows[0].cells[0].paragraphs[0].runs[0], sem_negrito=True)


def _aplica_formato(paragrafo, formato):
    if formato is None:
        return paragrafo
    for run in paragrafo.runs:
        atual = run._element.find(qn("w:rPr"))
        if atual is not None:
            run._element.remove(atual)
        run._element.insert(0, copy.deepcopy(formato))
    return paragrafo


def _larguras(tabela, polegadas):
    """Fixa a largura das colunas. Só mexer em tcW não basta: o LibreOffice segue o
    tblGrid, então as duas definições precisam concordar."""
    propriedades = tabela._tbl.find(qn("w:tblPr"))
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    propriedades.append(layout)

    twips = [int(round(valor * 1440)) for valor in polegadas]
    largura_total = OxmlElement("w:tblW")
    largura_total.set(qn("w:type"), "dxa")
    largura_total.set(qn("w:w"), str(sum(twips)))
    propriedades.append(largura_total)

    grade = tabela._tbl.find(qn("w:tblGrid"))
    for coluna, largura in zip(grade.findall(qn("w:gridCol")), twips):
        coluna.set(qn("w:w"), str(largura))

    tabela.autofit = False
    for linha in tabela.rows:
        for celula, largura in zip(linha.cells, polegadas):
            celula.width = Inches(largura)
    return tabela


def _ajusta_linhas(tabela, quantidade):
    """Deixa a tabela com o número de linhas de conteúdo necessário (fora o cabeçalho)."""
    modelo = copy.deepcopy(tabela.rows[-1]._tr)
    while len(tabela.rows) - 1 < quantidade:
        tabela._tbl.append(copy.deepcopy(modelo))
    while len(tabela.rows) - 1 > quantidade:
        tabela._tbl.remove(tabela.rows[-1]._tr)


def _limpa_fim(documento):
    """Descarta os parágrafos vazios do fim do template, que empurravam o documento
    para uma página em branco a mais."""
    for paragrafo in reversed(documento.paragraphs):
        if paragrafo.text.strip():
            break
        paragrafo._element.getparent().remove(paragrafo._element)


def main():
    documento = Document(TEMPLATE)
    # guardo a fonte do corpo antes de substituir os textos de exemplo, para reusá-la
    # nos parágrafos em branco do template, que não trazem formatação nenhuma
    formato_corpo = _formato(
        documento.paragraphs[_acha_indice(documento, "A empresa que você irá abordar") + 1].runs[0]
    )

    _troca(documento.paragraphs[0], "(TÍTULO DO SEU PROJETO)", TITULO)
    _troca(documento.paragraphs[3], "(NOME DA DISCIPLINA)", DISCIPLINA)
    _troca(documento.paragraphs[4], "(AUTORES)", AUTORES)

    _preenche_secao(documento, "A empresa que você irá abordar", EMPRESA)
    _preenche_secao(documento, "Área de negócio e o tema abordado", AREA_E_TEMA)

    perguntas, bases = documento.tables
    formato_perguntas = _formato_do_cabecalho(perguntas)
    _ajusta_linhas(perguntas, len(PERGUNTAS))
    for linha, pergunta in zip(perguntas.rows[1:], PERGUNTAS):
        celula = _texto(linha.cells[0].paragraphs[0], pergunta)
        celula.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _aplica_formato(celula, formato_perguntas)

    intro = _texto(documento.paragraphs[_acha_indice(documento, "Quais bases de dados") + 1],
                   INTRO_BASES)
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _aplica_formato(intro, formato_corpo)

    formato_bases = _formato_do_cabecalho(bases)
    _ajusta_linhas(bases, len(BASES))
    _larguras(bases, (2.85, 3.45))
    for linha, (base, link) in zip(bases.rows[1:], BASES):
        celula = _texto(linha.cells[0].paragraphs[0], base)
        celula.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _aplica_formato(celula, formato_bases)
        _aplica_formato(_texto(linha.cells[1].paragraphs[0], link), formato_bases)

    _limpa_fim(documento)

    SAIDA.mkdir(parents=True, exist_ok=True)
    destino_docx = SAIDA / f"{ARQUIVO}.docx"
    destino_pdf = SAIDA / f"{ARQUIVO}.pdf"
    documento.save(destino_docx)
    _para_pdf(destino_docx, destino_pdf)
    print(f"  {destino_docx.relative_to(RAIZ)}")
    print(f"  {destino_pdf.relative_to(RAIZ)}")


if __name__ == "__main__":
    main()
